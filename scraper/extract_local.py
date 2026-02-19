#!/usr/bin/env python3
"""
Local document extraction - runs on Mac in parallel with EC2.
Uses M1 Max for Tesseract, Gemini API for item extraction.
"""
import os
import sys
import asyncio
import time
from pathlib import Path
from typing import Optional, List, Dict
from dotenv import load_dotenv

# Load .env
load_dotenv(Path(__file__).parent.parent / '.env')

import asyncpg
import aiohttp
import google.generativeai as genai

# Database (RDS - accessible from anywhere)
DATABASE_URL = os.getenv('DATABASE_URL')

# Gemini
GEMINI_API_KEY = os.getenv('GEMINI_API_KEY')
if GEMINI_API_KEY:
    genai.configure(api_key=GEMINI_API_KEY)

# Local dirs
DOWNLOAD_DIR = Path('/tmp/local_extraction')
DOWNLOAD_DIR.mkdir(exist_ok=True)


class LocalExtractor:
    def __init__(self):
        self.pool = None  # Use pool instead of single connection
        self.session = None
        self.gemini_model = None
        self.stats = {'processed': 0, 'items': 0, 'failed': 0}
        self.start_time = None

    async def connect(self):
        self.pool = await asyncpg.create_pool(DATABASE_URL, min_size=1, max_size=5)
        self.session = aiohttp.ClientSession()
        if GEMINI_API_KEY:
            self.gemini_model = genai.GenerativeModel('gemini-2.0-flash-exp')
        print(f"✓ Connected to database (pool)")
        print(f"✓ Gemini: {'available' if self.gemini_model else 'NOT available'}")

    async def close(self):
        if self.session:
            await self.session.close()
        if self.pool:
            await self.pool.close()

    async def get_pending_docs(self, limit: int, offset: int = 0) -> List[Dict]:
        """Get pending documents - use offset to avoid collision with EC2."""
        async with self.pool.acquire() as conn:
            rows = await conn.fetch("""
                SELECT doc_id::text, tender_id, file_url, file_name,
                       COALESCE(doc_type, 'other') as category
                FROM documents
                WHERE extraction_status = 'pending'
                AND file_url IS NOT NULL
                ORDER BY uploaded_at ASC
                LIMIT $1 OFFSET $2
            """, limit, offset)
            return [dict(r) for r in rows]

    async def download_doc(self, doc: Dict) -> Optional[Path]:
        """Download document."""
        safe_name = f"{doc['tender_id'].replace('/', '_')}_{hash(doc['file_url']) & 0xFFFF:04x}"
        ext = Path(doc['file_name']).suffix.lower() or '.pdf'
        local_path = DOWNLOAD_DIR / f"{safe_name}{ext}"

        if local_path.exists():
            return local_path

        try:
            async with self.session.get(doc['file_url'], timeout=aiohttp.ClientTimeout(total=30)) as resp:
                if resp.status == 200:
                    content = await resp.read()
                    local_path.write_bytes(content)
                    return local_path
        except Exception as e:
            print(f"  ✗ Download failed: {e}")
        return None

    def extract_text(self, file_path: Path) -> Optional[str]:
        """Extract text from document."""
        import fitz  # PyMuPDF
        import pytesseract
        from PIL import Image
        import io

        ext = file_path.suffix.lower()

        try:
            if ext == '.pdf':
                doc = fitz.open(file_path)
                text_parts = []

                for page_num, page in enumerate(doc):
                    # Try text extraction first
                    text = page.get_text()
                    if len(text.strip()) > 100:
                        text_parts.append(text)
                    else:
                        # OCR for scanned pages
                        pix = page.get_pixmap(dpi=150)
                        img = Image.open(io.BytesIO(pix.tobytes("png")))
                        ocr_text = pytesseract.image_to_string(img, lang='mkd+eng')
                        text_parts.append(ocr_text)

                doc.close()
                return '\n\n'.join(text_parts)

            elif ext in ['.docx', '.doc']:
                from docx import Document
                doc = Document(file_path)
                return '\n'.join([p.text for p in doc.paragraphs])

            elif ext in ['.xlsx', '.xls']:
                import openpyxl
                wb = openpyxl.load_workbook(file_path, data_only=True)
                text_parts = []
                for sheet in wb.worksheets:
                    for row in sheet.iter_rows(values_only=True):
                        text_parts.append(' | '.join(str(c) if c else '' for c in row))
                return '\n'.join(text_parts)

        except Exception as e:
            print(f"  ✗ Extraction error: {e}")
        return None

    async def extract_items_gemini(self, text: str, category: str) -> List[Dict]:
        """Use Gemini to extract items."""
        if not self.gemini_model or len(text) < 100:
            return []

        prompt = f"""Extract product items from this {category} document.
Return JSON array with objects containing:
- name: product name
- quantity: number
- unit: unit of measure
- unit_price: price per unit
- total_price: total price

Document text:
{text[:15000]}

Return ONLY valid JSON array, no markdown."""

        try:
            response = await asyncio.to_thread(
                self.gemini_model.generate_content, prompt
            )
            import json
            result = response.text.strip()
            if result.startswith('```'):
                result = result.split('\n', 1)[1].rsplit('```', 1)[0]
            return json.loads(result)
        except Exception as e:
            return []

    def parse_decimal(self, value):
        """Parse decimal value handling European comma format."""
        if value is None:
            return None
        if isinstance(value, (int, float)):
            return float(value)
        if isinstance(value, str):
            # Handle European format: 58,47 -> 58.47
            value = value.replace(',', '.').replace(' ', '')
            try:
                return float(value)
            except:
                return None
        return None

    async def save_results(self, doc_id: str, tender_id: str, text: str, items: List[Dict], category: str):
        """Save extraction results."""
        import json

        # Strip null bytes that cause PostgreSQL UTF-8 errors
        clean_text = text.replace('\x00', '') if text else text

        async with self.pool.acquire() as conn:
            # Update document
            await conn.execute("""
                UPDATE documents SET
                    content_text = $2,
                    extraction_status = 'success',
                    extraction_method = 'local_gemini'
                WHERE doc_id = $1::uuid
            """, doc_id, clean_text)

            # Insert items
            for i, item in enumerate(items):
                name = item.get('name', '').strip()
                if not name or len(name) < 3:
                    continue

                raw_json = json.dumps(item, ensure_ascii=False, default=str)

                await conn.execute("""
                    INSERT INTO product_items (
                        tender_id, document_id, item_number,
                        name, quantity, unit, unit_price, total_price,
                        extraction_method, raw_text
                    ) VALUES ($1, $2::uuid, $3, $4, $5, $6, $7, $8, $9, $10)
                    ON CONFLICT DO NOTHING
                """,
                    tender_id, doc_id, i + 1, name,
                    self.parse_decimal(item.get('quantity')), item.get('unit'),
                    self.parse_decimal(item.get('unit_price')), self.parse_decimal(item.get('total_price')),
                    f'local_gemini_{category}', raw_json
                )
                self.stats['items'] += 1

    async def mark_failed(self, doc_id: str, status: str):
        """Mark document as failed so we don't retry it."""
        async with self.pool.acquire() as conn:
            await conn.execute("""
                UPDATE documents SET extraction_status = $2
                WHERE doc_id = $1::uuid
            """, doc_id, status)

    async def process_doc(self, doc: Dict) -> bool:
        """Process a single document."""
        doc_id = doc['doc_id']
        tender_id = doc['tender_id']
        category = doc['category']

        # Download
        file_path = await self.download_doc(doc)
        if not file_path:
            await self.mark_failed(doc_id, 'download_failed')
            return False

        # Extract text
        text = self.extract_text(file_path)
        if not text or len(text) < 50:
            file_path.unlink(missing_ok=True)
            await self.mark_failed(doc_id, 'skip_empty')
            return False

        # Gemini extraction
        items = await self.extract_items_gemini(text, category)

        # Save
        await self.save_results(doc_id, tender_id, text, items, category)

        # Cleanup
        file_path.unlink(missing_ok=True)

        self.stats['processed'] += 1
        return True

    async def run(self, limit: int = 100, offset: int = 0, concurrency: int = 3):
        """Run extraction with concurrency."""
        self.start_time = time.time()
        await self.connect()

        try:
            docs = await self.get_pending_docs(limit, offset)
            print(f"\n📄 Found {len(docs)} documents (offset={offset})\n")

            if not docs:
                return

            sem = asyncio.Semaphore(concurrency)

            async def process_with_sem(doc):
                async with sem:
                    fname = doc['file_name'][:40]
                    try:
                        ok = await self.process_doc(doc)
                        status = "✓" if ok else "✗"
                        print(f"  {status} {fname}")
                    except Exception as e:
                        print(f"  ✗ {fname}: {e}")
                        self.stats['failed'] += 1

            await asyncio.gather(*[process_with_sem(d) for d in docs])

            elapsed = time.time() - self.start_time
            rate = self.stats['processed'] / elapsed * 60 if elapsed > 0 else 0

            print(f"\n{'='*50}")
            print(f"Completed in {elapsed:.1f}s ({rate:.1f}/min)")
            print(f"Processed: {self.stats['processed']}, Items: {self.stats['items']}, Failed: {self.stats['failed']}")

        finally:
            await self.close()


async def main():
    import argparse
    parser = argparse.ArgumentParser()
    parser.add_argument('--limit', type=int, default=100)
    parser.add_argument('--offset', type=int, default=0, help='Offset to avoid collision with EC2')
    parser.add_argument('--concurrency', type=int, default=3)
    args = parser.parse_args()

    print("="*50)
    print("Local Extraction (Mac M1 + Gemini)")
    print("="*50)

    extractor = LocalExtractor()
    await extractor.run(args.limit, args.offset, args.concurrency)


if __name__ == '__main__':
    asyncio.run(main())
